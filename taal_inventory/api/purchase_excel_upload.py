"""
taal_inventory/api/purchase_excel_upload.py

Supplier Excel -> Draft Purchase Invoice (button on Purchase Invoice)
=====================================================================
No separate upload DocType. A button on the Purchase Invoice list (and
optionally on the form) opens a dialog: attach the Tally invoice Excel
(layout of Purchase_20.xlsx / Purchase_19.xlsx), press Create Draft.
Works for ANY supplier: the Supplier is matched from the file by GSTIN,
then by name, and the GST template (In-state / Out-state) is picked from
the two GSTIN state codes on the file. Both can be overridden in the
dialog, and the result table shows what was used.

This method parses the file, validates every row, creates the Purchase
Invoice as a DRAFT, and attaches the Excel to that invoice for audit.
It never submits. You review the draft, fill what is marked, and submit
yourself. Your events/purchase_invoice.py on_submit then creates or
tops up the Pharmacy Batches as usual.

Guarantees, unchanged from before:
- Reuses _get_pharmacy_item and _find_batch from events/purchase_invoice.py,
  so CREATE vs TOPUP is decided here EXACTLY as on_submit will decide it
  later (Active/Exhausted only, company scoped, newest first).
- Every row must resolve to an Item that has a Pharmacy Item record. This
  matches the pharmacy-only validate() guard.
- Existing batch (top-up): MRP, selling price, expiry and mfg date are
  copied FROM the batch, mirroring the Purchase Invoice client script.
  Conflicts between the Excel and the batch are flagged in the result
  table, never silently fixed.
- New batch: Excel expiry is used. MRP and selling price stay blank, fill
  them in the draft. _validate_pharmacy_row blocks submit until filled.

Location : apps/taal_inventory/taal_inventory/api/purchase_excel_upload.py
Method   : taal_inventory.api.purchase_excel_upload.create_purchase_invoice
After copying: bench --site <site> clear-cache && bench restart
"""

import io
import re
from datetime import date, datetime

import frappe
import openpyxl
from frappe import _
from frappe.utils import cint, flt, getdate

from taal_inventory.events.purchase_invoice import _find_batch, _get_pharmacy_item

DATE_FORMATS = ("%d-%b-%y", "%d-%b-%Y", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d")


# ------------------------------------------------------------------ helpers

def clean(value):
    return "" if value is None else str(value).strip()


def as_serial(value):
    """Return a positive int if the cell holds an item serial number."""
    try:
        number = int(float(clean(value)))
        return number if number > 0 else None
    except (TypeError, ValueError):
        return None


def as_number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_date_value(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def label_value(row, prefix):
    """First non-empty cell after a cell whose text starts with prefix."""
    for index, cell in enumerate(row):
        text = clean(cell).lower().rstrip(":").strip()
        if text.startswith(prefix):
            for value in row[index + 1:]:
                if value not in (None, ""):
                    return value
            return None
    return None


def row_text(row):
    return " ".join(clean(cell).lower() for cell in row if clean(cell))


def squash(text):
    return re.sub(r"\s+", " ", clean(text)).lower()


# ------------------------------------------------------------------ parsing

def find_item_columns(grid):
    """Locate the item table header row and its column positions."""
    for index, row in enumerate(grid):
        labels = {}
        for col, cell in enumerate(row):
            text = clean(cell).lower().rstrip(".")
            if text:
                labels[text] = col
        name_col = next((labels[key] for key in labels if key.startswith("description")), None)
        if name_col is not None and "quantity" in labels and "amount" in labels:
            return index, {
                "serial": labels.get("sl", 0),
                "name": name_col,
                "qty": labels["quantity"],
                "rate": labels.get("rate"),
                "per": labels.get("per"),
                "amount": labels["amount"],
            }
    return None, None


GSTIN_RE = re.compile(r"\b\d{2}[A-Z]{5}\d{4}[A-Z][1-9A-Z]Z[0-9A-Z]\b")


def read_header_block(grid, header_index, data):
    """
    Invoice no, dates, supplier name and both GSTINs from the rows above the
    item table. The buyer (your company) block comes first, then the
    "Supplier (Bill from)" block. So the first GSTIN after that label is the
    supplier's, and the first GSTIN before it is the buyer's.
    """
    supplier_row = None
    gstins = []  # (row_index, gstin) in order of appearance

    for index in range(header_index):
        row = grid[index]
        next_row = grid[index + 1] if index + 1 < len(grid) else []
        for col, cell in enumerate(row):
            text = clean(cell)
            low = text.lower()
            for found in GSTIN_RE.findall(text.upper()):
                gstins.append((index, found))
            if low.startswith("invoice no") and col < len(next_row):
                data["invoice_no"] = clean(next_row[col])
            elif low == "dated" and col < len(next_row):
                data["invoice_date"] = parse_date_value(next_row[col])
            elif low.startswith("supplier invoice no") and col < len(next_row):
                raw = clean(next_row[col])
                match = re.match(r"^(\S+)\s+dt\.?\s*(.+)$", raw)
                if match:
                    data["ref_no"] = match.group(1)
                    data["ref_date"] = parse_date_value(match.group(2))
                elif raw:
                    data["ref_no"] = raw
            elif low.startswith("supplier (bill from)") and next_row:
                data["supplier_name"] = clean(next_row[0])
                supplier_row = index

    if supplier_row is not None:
        after = [g for r, g in gstins if r > supplier_row]
        before = [g for r, g in gstins if r < supplier_row]
        data["supplier_gstin"] = after[0] if after else ""
        data["buyer_gstin"] = before[0] if before else ""


def read_item_block(grid, header_index, columns, data):
    """Item rows (3 lines each: item / Batch / Expiry) and the totals block."""

    def cell(row, key):
        col = columns.get(key)
        return row[col] if col is not None and col < len(row) else None

    current = None
    for index in range(header_index + 1, len(grid)):
        row = grid[index]
        serial = as_serial(cell(row, "serial"))
        name = clean(cell(row, "name"))
        text = row_text(row)
        amount = as_number(cell(row, "amount"))

        if serial and name:
            current = {
                "row": index + 1,
                "name": name,
                "qty": flt(cell(row, "qty")),
                "rate": flt(cell(row, "rate")),
                "uom": clean(cell(row, "per")),
                "amount": flt(amount),
                "batch": "",
                "expiry": None,
                "expiry_raw": "",
            }
            data["items"].append(current)
        elif name.lower() == "total":
            data["total_qty"] = flt(cell(row, "qty"))
            data["grand_total"] = flt(amount)
            break
        elif "cgst" in text:
            data["cgst"] = flt(amount)
        elif "sgst" in text:
            data["sgst"] = flt(amount)
        elif "igst" in text:
            data["igst"] = flt(amount)
        elif "round" in text:
            data["round_off"] = flt(amount)
        elif current and label_value(row, "batch") is not None:
            current["batch"] = clean(label_value(row, "batch"))
        elif current and label_value(row, "exp") is not None:
            raw = label_value(row, "exp")
            current["expiry_raw"] = clean(raw)
            current["expiry"] = parse_date_value(raw)
        elif data["items"] and amount is not None and data["taxable_total"] is None and not name:
            data["taxable_total"] = flt(amount)


def parse_workbook(content):
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheet = None
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().lower() == "invoice":
            sheet = workbook[sheet_name]
            break
    sheet = sheet or workbook.active
    grid = [list(row) for row in sheet.iter_rows(values_only=True)]

    data = {
        "supplier_name": "",
        "supplier_gstin": "",
        "buyer_gstin": "",
        "invoice_no": "",
        "invoice_date": None,
        "ref_no": "",
        "ref_date": None,
        "items": [],
        "taxable_total": None,
        "cgst": None,
        "sgst": None,
        "igst": None,
        "round_off": None,
        "grand_total": None,
        "total_qty": None,
    }

    header_index, columns = find_item_columns(grid)
    if header_index is None:
        frappe.throw(_(
            "Could not find the item table header "
            "(Description of Goods / Quantity / Rate / Amount) in this file. "
            "Is this the usual Tally invoice export?"))

    read_header_block(grid, header_index, data)
    read_item_block(grid, header_index, columns, data)
    return data


# --------------------------------------------------------------- validation

def resolve_item(name, supplier):
    """
    Collect EVERY Item the file name could mean, in this order:
      1. Item Code equals the name
      2. Item Name equals the name
      3. Pharmacy Item Product Name equals the name -> its Item Code
      4. Item Supplier -> Supplier Part Number
    Then return the FIRST candidate that has a Pharmacy Item record.

    Why: two Items can share a name. Example: an Item whose Item Code is
    literally "Spegra" next to TPIC 67 whose Item Name is "Spegra". Only
    TPIC 67 is a pharmacy item, so TPIC 67 must win regardless of order.
    If no candidate has a Pharmacy Item, the first one is returned so the
    error message can name it.
    """
    candidates = []

    def add(code):
        if code and code not in candidates:
            candidates.append(code)

    add(frappe.db.get_value("Item", {"name": name, "disabled": 0}))

    for code in frappe.get_all("Item", filters={"item_name": name, "disabled": 0}, pluck="name"):
        add(code)

    for code in frappe.get_all("Pharmacy Item", filters={"item_name": name}, pluck="item_code"):
        if code and frappe.db.exists("Item", {"name": code, "disabled": 0}):
            add(code)

    add(frappe.db.get_value(
        "Item Supplier",
        {"supplier": supplier, "supplier_part_no": name},
        "parent",
    ))

    for code in candidates:
        if _get_pharmacy_item(code):
            return code
    return candidates[0] if candidates else None


def merge_duplicate_rows(items):
    """
    Same item + same batch + same rate on several lines -> one line
    (your _validate_pharmacy_row requires one line per batch at submit).
    Same batch with different rates, or the same batch under two different
    items, is an error.
    """
    errors = []
    merged = []
    index = {}

    for entry in items:
        key = (entry.get("item_code"), entry.get("batch"))
        if entry.get("item_code") and entry.get("batch") and key in index:
            first = index[key]
            if abs(flt(first["rate"]) - flt(entry["rate"])) <= 0.01:
                first["qty"] = flt(first["qty"] + entry["qty"])
                first["amount"] = flt(first["amount"] + entry["amount"], 2)
                continue
            errors.append(_(
                "Rows {0} and {1} ({2}): same batch {3} with two different "
                "rates. One batch must have one rate, fix the file.").format(
                first["row"], entry["row"], entry["name"], entry["batch"]))
        index[key] = entry
        merged.append(entry)

    batch_owners = {}
    for entry in merged:
        if entry.get("batch") and entry.get("item_code"):
            batch_owners.setdefault(entry["batch"], set()).add(entry["item_code"])
    for batch_id, owners in batch_owners.items():
        if len(owners) > 1:
            errors.append(_(
                "Batch {0} appears under more than one item ({1}). "
                "Submit would block this, fix the file.").format(
                batch_id, ", ".join(sorted(owners))))

    return merged, errors


def find_supplier(data):
    """
    Resolve the ERPNext Supplier from the file.
      1. GSTIN on the Supplier record (India Compliance field)
      2. GSTIN on an Address linked to exactly one Supplier
      3. Supplier ID or Supplier Name equal to the "Supplier (Bill from)" text
    Returns the Supplier name, or None when nothing matches.
    """
    gstin = data.get("supplier_gstin") or ""

    if gstin:
        if frappe.get_meta("Supplier").has_field("gstin"):
            found = frappe.db.get_value("Supplier", {"gstin": gstin, "disabled": 0})
            if found:
                return found
        if frappe.get_meta("Address").has_field("gstin"):
            addresses = frappe.get_all("Address", filters={"gstin": gstin}, pluck="name")
            if addresses:
                linked = frappe.get_all(
                    "Dynamic Link",
                    filters={"parenttype": "Address", "parent": ["in", addresses],
                             "link_doctype": "Supplier"},
                    pluck="link_name")
                linked = [s for s in dict.fromkeys(linked)
                          if not frappe.db.get_value("Supplier", s, "disabled")]
                if len(linked) == 1:
                    return linked[0]

    wanted = squash(data.get("supplier_name"))
    if wanted:
        for row in frappe.get_all("Supplier", filters={"disabled": 0},
                                  fields=["name", "supplier_name"]):
            if squash(row.name) == wanted or squash(row.supplier_name) == wanted:
                return row.name
    return None


STATE_BY_CODE = {
    "01": "Jammu and Kashmir", "02": "Himachal Pradesh", "03": "Punjab", "04": "Chandigarh",
    "05": "Uttarakhand", "06": "Haryana", "07": "Delhi", "08": "Rajasthan", "09": "Uttar Pradesh",
    "10": "Bihar", "11": "Sikkim", "12": "Arunachal Pradesh", "13": "Nagaland", "14": "Manipur",
    "15": "Mizoram", "16": "Tripura", "17": "Meghalaya", "18": "Assam", "19": "West Bengal",
    "20": "Jharkhand", "21": "Odisha", "22": "Chhattisgarh", "23": "Madhya Pradesh", "24": "Gujarat",
    "26": "Dadra and Nagar Haveli and Daman and Diu", "27": "Maharashtra", "29": "Karnataka",
    "30": "Goa", "31": "Lakshadweep Islands", "32": "Kerala", "33": "Tamil Nadu", "34": "Puducherry",
    "35": "Andaman and Nicobar Islands", "36": "Telangana", "37": "Andhra Pradesh", "38": "Ladakh",
}


def create_supplier_from_file(data):
    """
    Creates a Supplier from the "Supplier (Bill from)" block, plus a billing
    Address when a GSTIN is on the file. Only called when the dialog checkbox
    "Create supplier if not found" was ticked. Runs inside the same
    transaction as the invoice, so if anything later fails nothing stays.
    The result table flags the new supplier for review before submit.
    """
    name = clean(data.get("supplier_name"))
    gstin = data.get("supplier_gstin") or ""
    if not name:
        frappe.throw(_("The file has no supplier name, so the supplier cannot be created."))

    supplier_meta = frappe.get_meta("Supplier")
    supplier = frappe.new_doc("Supplier")
    supplier.supplier_name = name
    supplier.supplier_group = "All Supplier Groups"
    supplier.supplier_type = "Company"
    if supplier_meta.has_field("country"):
        supplier.country = "India"
    if supplier_meta.has_field("gst_category"):
        supplier.gst_category = "Registered Regular" if gstin else "Unregistered"
    if gstin and supplier_meta.has_field("gstin"):
        supplier.gstin = gstin
    supplier.insert()

    data["supplier_created"] = supplier.name
    data["address_created"] = ""

    if gstin:
        state = STATE_BY_CODE.get(gstin[:2], "")
        try:
            address_meta = frappe.get_meta("Address")
            address = frappe.new_doc("Address")
            address.address_title = name
            address.address_type = "Billing"
            address.address_line1 = _("Address to be updated (auto-created from supplier Excel)")
            address.city = state or _("Unknown")
            address.state = state
            address.country = "India"
            address.is_primary_address = 1
            if address_meta.has_field("gstin"):
                address.gstin = gstin
            if address_meta.has_field("gst_category"):
                address.gst_category = "Registered Regular"
            if address_meta.has_field("gst_state") and state:
                address.gst_state = state
            address.append("links", {"link_doctype": "Supplier", "link_name": supplier.name})
            address.insert()
            data["address_created"] = address.name
        except Exception:
            frappe.log_error(frappe.get_traceback(), "Purchase Excel Upload supplier address")

    return supplier.name


def pick_tax_template(company, data):
    """
    In-state when the supplier GSTIN state code equals the buyer GSTIN state
    code on the file, out-state otherwise. Without both GSTINs, an IGST line
    on the file means out-state. Returns the template name, or None if that
    template does not exist for the company.
    """
    supplier_state = (data.get("supplier_gstin") or "")[:2]
    buyer_state = (data.get("buyer_gstin") or "")[:2]
    if supplier_state and buyer_state:
        in_state = supplier_state == buyer_state
    elif data.get("igst"):
        in_state = False
    else:
        in_state = True

    abbr = frappe.db.get_value("Company", company, "abbr")
    template = "Input GST {0} - {1}".format("In-state" if in_state else "Out-state", abbr)
    if frappe.db.exists("Purchase Taxes and Charges Template", template):
        return template
    return None


def validate_parsed(args, data):
    errors = []

    # Supplier: the dialog choice wins. Blank means take it from the file.
    args.supplier_chosen = bool(args.supplier)
    if not args.supplier:
        args.supplier = find_supplier(data) or ""
        if not args.supplier and cint(args.get("create_supplier")):
            args.supplier = create_supplier_from_file(data)
        if not args.supplier:
            errors.append(_(
                "Supplier {0} (GSTIN {1}) was not found in the Supplier master. "
                "Tick 'Create supplier if not found' in the dialog, create it "
                "yourself with this GSTIN, or pick an existing supplier.").format(
                frappe.bold(data["supplier_name"] or _("unknown")),
                data["supplier_gstin"] or _("not on file")))

    # Tax template: the dialog choice wins. Blank means pick by GSTIN state.
    if not args.taxes_and_charges:
        args.taxes_and_charges = pick_tax_template(args.company, data) or ""
        if not args.taxes_and_charges:
            errors.append(_(
                "Could not pick a GST template automatically (In-state / Out-state "
                "template not found for this company). Choose one in the dialog."))

    if not data["items"]:
        errors.append(_("No item rows were found in the file."))

    # Resolve every row to an Item and its Pharmacy Item first.
    for entry in data["items"]:
        entry["item_code"] = resolve_item(entry["name"], args.supplier)
        entry["pharmacy_item"] = _get_pharmacy_item(entry["item_code"]) if entry["item_code"] else None

    merged, merge_errors = merge_duplicate_rows(data["items"])
    data["items"] = merged
    errors.extend(merge_errors)

    for entry in data["items"]:
        prefix = _("Row {0} ({1}):").format(entry["row"], entry["name"])

        if not entry["item_code"]:
            errors.append(prefix + " " + _(
                "no matching Item found. Match by Item Code, Item Name, "
                "Pharmacy Item Product Name, or add this exact text as "
                "Supplier Part Number in the Item's Supplier Items table."))
        elif not entry["pharmacy_item"]:
            errors.append(prefix + " " + _(
                "matched Item {0}, which has no Pharmacy Item record. "
                "Purchases are pharmacy items only. Either link this Item "
                "to a Pharmacy Item, or if it is a duplicate of another "
                "Item with the same name, disable it.").format(
                frappe.bold(entry["item_code"])))

        if entry["qty"] <= 0:
            errors.append(prefix + " " + _("quantity is missing or zero."))
        if entry["rate"] <= 0:
            errors.append(prefix + " " + _("rate is missing or zero."))
        if abs(flt(entry["qty"] * entry["rate"], 2) - flt(entry["amount"], 2)) > 0.05:
            errors.append(prefix + " " + _("quantity x rate does not match the row amount in the file."))
        if not entry["batch"]:
            errors.append(prefix + " " + _("batch number is missing."))

        # Existing batch lookup, IDENTICAL to what on_submit will decide.
        entry["existing_batch"] = None
        entry["batch_expiry"] = None
        entry["batch_mfg"] = None
        if entry["pharmacy_item"] and entry["batch"]:
            existing = _find_batch(entry["batch"], entry["pharmacy_item"], args.company)
            if existing:
                entry["existing_batch"] = existing
                extra = frappe.db.get_value(
                    "Pharmacy Batch", existing.name,
                    ["expiry_date", "manufacturing_date"], as_dict=True) or {}
                entry["batch_expiry"] = extra.get("expiry_date")
                entry["batch_mfg"] = extra.get("manufacturing_date")

        # New batch needs an expiry from the file. Top-up takes the batch's.
        if not entry["expiry"] and not entry["existing_batch"]:
            errors.append(prefix + " " + _("expiry date is missing or unreadable ({0}).").format(
                entry["expiry_raw"] or _("blank")))

    bill_no = clean(args.get("bill_no")) or data["ref_no"] or data["invoice_no"]
    bill_date = args.get("bill_date") or data["ref_date"] or data["invoice_date"]
    if not bill_no:
        errors.append(_("No bill number found in the file. Enter Bill No in the dialog and retry."))
    if not bill_date:
        errors.append(_("No bill date found in the file. Enter Bill Date in the dialog and retry."))
    data["bill_no"] = bill_no
    data["bill_date"] = bill_date

    if bill_no and args.supplier:
        existing = frappe.db.get_value("Purchase Invoice", {
            "supplier": args.supplier,
            "bill_no": bill_no,
            "docstatus": ("<", 2),
        })
        if existing:
            errors.append(_("Bill No {0} is already entered as {1}. Nothing was created.").format(
                frappe.bold(bill_no), frappe.bold(existing)))

    return errors


# ----------------------------------------------------------------- creation

def build_invoice(args, data):
    item_meta = frappe.get_meta("Purchase Invoice Item")

    def set_field(row, field, value):
        if value not in (None, "") and item_meta.has_field(field):
            row.set(field, value)

    invoice = frappe.new_doc("Purchase Invoice")
    invoice.company = args.company
    invoice.supplier = args.supplier
    invoice.bill_no = data["bill_no"]
    invoice.bill_date = data["bill_date"]
    invoice.set_posting_time = 1
    invoice.posting_date = args.get("posting_date") or data["bill_date"]
    invoice.due_date = invoice.posting_date
    invoice.update_stock = 0
    invoice.ignore_pricing_rule = 1
    if args.get("excel_file") and frappe.get_meta("Purchase Invoice").has_field("custom_supplier_excel"):
        invoice.custom_supplier_excel = args.excel_file  # drives the View Supplier Excel button

    for entry in data["items"]:
        row = invoice.append("items", {
            "item_code": entry["item_code"],
            "qty": entry["qty"],
            "rate": entry["rate"],
        })
        if args.get("set_warehouse"):
            row.warehouse = args.set_warehouse

        stock_uom = frappe.db.get_value("Item", entry["item_code"], "stock_uom")
        if entry["uom"] and entry["uom"] == stock_uom:
            row.uom = entry["uom"]
            row.conversion_factor = 1

        set_field(row, "custom_batch_id", entry["batch"])

        if entry.get("existing_batch"):
            # TOPUP: mirror the batch, same as your PI client script does.
            existing = entry["existing_batch"]
            set_field(row, "custom_expiry_date", entry["batch_expiry"] or entry["expiry"])
            set_field(row, "custom_manufacturing_date", entry["batch_mfg"])
            set_field(row, "custom_mrp_price", flt(existing.mrp_price))
            set_field(row, "custom_selling_price", flt(existing.selling_price))
        else:
            # CREATE: expiry from the file. MRP and Selling Price are not in
            # the Tally file. Fill them in the draft, submit is blocked by
            # _validate_pharmacy_row until they are filled.
            set_field(row, "custom_expiry_date", entry["expiry"])

    if args.get("taxes_and_charges"):
        invoice.taxes_and_charges = args.taxes_and_charges
        template = frappe.get_doc("Purchase Taxes and Charges Template", args.taxes_and_charges)
        for tax in template.taxes:
            invoice.append("taxes", {
                "category": tax.category,
                "add_deduct_tax": tax.add_deduct_tax,
                "charge_type": tax.charge_type,
                "account_head": tax.account_head,
                "description": tax.description,
                "rate": tax.rate,
                "cost_center": tax.cost_center,
                "included_in_print_rate": tax.included_in_print_rate,
            })

    invoice.insert()  # docstatus 0 = Draft. Submission stays manual.
    return invoice


def verify_invoice(invoice, data):
    """Compare the draft against the file. Fix rates if a Price List overrode them."""
    corrections = False
    for row, entry in zip(invoice.items, data["items"]):
        if flt(row.rate, 2) != flt(entry["rate"], 2) or flt(row.amount, 2) != flt(entry["amount"], 2):
            if flt(row.price_list_rate) > 0 and flt(row.price_list_rate) >= flt(entry["rate"]):
                row.margin_type = ""
                row.margin_rate_or_amount = 0
                row.discount_percentage = 0
                row.discount_amount = flt(row.price_list_rate - entry["rate"], 2)
            else:
                row.price_list_rate = 0
                row.discount_percentage = 0
                row.discount_amount = 0
            row.rate = flt(entry["rate"])
            corrections = True
    if corrections:
        invoice.save()
        invoice.load_from_db()

    summary = []

    def fmt(value):
        return f"{flt(value):,.2f}"

    def add(label, excel_value, erp_value, tolerance):
        summary.append({
            "label": label,
            "excel": fmt(excel_value),
            "erp": fmt(erp_value),
            "ok": abs(flt(excel_value) - flt(erp_value)) <= tolerance,
        })

    summary.append({
        "label": _("Bill No / Bill Date"),
        "excel": "{0} / {1}".format(data["bill_no"], data["bill_date"]),
        "erp": "{0} / {1}".format(invoice.bill_no, invoice.bill_date),
        "ok": True,
    })
    file_supplier = data.get("supplier_name") or ""
    file_gstin = data.get("supplier_gstin") or ""
    invoice_supplier = frappe.db.get_value("Supplier", invoice.supplier, "supplier_name") or invoice.supplier
    invoice_gstin = invoice.get("supplier_gstin") or ""
    supplier_ok = bool(
        (file_gstin and file_gstin == invoice_gstin)
        or squash(file_supplier) == squash(invoice_supplier)
    )
    summary.append({
        "label": _("Supplier (file / invoice)"),
        "excel": "{0} {1}".format(file_supplier, file_gstin).strip(),
        "erp": "{0} {1}".format(invoice_supplier, invoice_gstin).strip(),
        "ok": supplier_ok,
    })
    if data.get("supplier_created"):
        summary.append({
            "label": _("NEW supplier created from the file. Open it, complete the "
                       "address and GST details, then come back and submit"),
            "excel": data.get("supplier_created"),
            "erp": data.get("address_created") or _("no address (GSTIN missing or address failed)"),
            "ok": False,
        })
    add(_("Item lines"), len(data["items"]), len(invoice.items), 0)
    if data["total_qty"]:
        add(_("Total quantity"), data["total_qty"], invoice.total_qty, 0)
    if data["taxable_total"]:
        add(_("Net total (before GST)"), data["taxable_total"], invoice.net_total, 0.01)

    cgst = sgst = igst = 0
    for tax in invoice.taxes:
        text = (tax.description or "").lower()
        if "cgst" in text:
            cgst = tax.tax_amount
        elif "sgst" in text:
            sgst = tax.tax_amount
        elif "igst" in text:
            igst = tax.tax_amount
    if data["cgst"] is not None:
        add(_("CGST"), data["cgst"], cgst, 0.05)
    if data["sgst"] is not None:
        add(_("SGST"), data["sgst"], sgst, 0.05)
    if data["igst"] is not None:
        add(_("IGST"), data["igst"], igst, 0.05)

    file_uses_igst = bool(data.get("igst"))
    invoice_uses_igst = any("igst" in (t.description or "").lower() for t in invoice.taxes)
    if file_uses_igst != invoice_uses_igst:
        summary.append({
            "label": _("GST type on file vs template {0}").format(
                invoice.taxes_and_charges or _("(none)")),
            "excel": "IGST" if file_uses_igst else "CGST + SGST",
            "erp": "IGST" if invoice_uses_igst else "CGST + SGST",
            "ok": False,
        })
    if data["grand_total"]:
        add(_("Grand total (rounded)"), data["grand_total"],
            invoice.rounded_total or invoice.grand_total, 1.0)

    # Batch picture, so you know what submit will do.
    topups = [e for e in data["items"] if e.get("existing_batch")]
    new_batches = [e for e in data["items"] if not e.get("existing_batch")]

    if topups:
        summary.append({
            "label": _("Existing batches (top-up, pricing copied from batch)"),
            "excel": str(len(topups)),
            "erp": str(len(topups)),
            "ok": True,
        })
    if new_batches:
        summary.append({
            "label": _("New batches, fill MRP and Selling Price in the draft "
                       "before submit: {0}").format(
                ", ".join(e["name"] for e in new_batches)),
            "excel": str(len(new_batches)),
            "erp": str(len(new_batches)),
            "ok": False,
        })

    for entry in topups:
        if entry["expiry"] and entry["batch_expiry"] and entry["expiry"] != entry["batch_expiry"]:
            summary.append({
                "label": _("Expiry in file differs from batch {0} ({1}), "
                           "batch expiry kept").format(entry["batch"], entry["name"]),
                "excel": str(entry["expiry"]),
                "erp": str(entry["batch_expiry"]),
                "ok": False,
            })
        batch_cost = flt(entry["existing_batch"].cost_price)
        if batch_cost and abs(batch_cost - flt(entry["rate"])) > 0.01:
            summary.append({
                "label": _("Rate in file differs from batch cost price ({0}), "
                           "file rate kept").format(entry["name"]),
                "excel": fmt(entry["rate"]),
                "erp": fmt(batch_cost),
                "ok": False,
            })

    for row, entry in zip(invoice.items, data["items"]):
        if abs(flt(row.amount, 2) - flt(entry["amount"], 2)) > 0.01:
            summary.append({
                "label": _("Row amount: {0}").format(row.item_name or row.item_code),
                "excel": fmt(entry["amount"]),
                "erp": fmt(row.amount),
                "ok": False,
            })

    return summary


def _attach_file_to_invoice(file_doc, invoice):
    """Keep the source Excel on the invoice for audit. Never blocks creation."""
    try:
        if not file_doc.attached_to_name:
            frappe.db.set_value("File", file_doc.name, {
                "attached_to_doctype": "Purchase Invoice",
                "attached_to_name": invoice.name,
            }, update_modified=False)
    except Exception:
        frappe.log_error(frappe.get_traceback(), "Purchase Excel Upload attach")


# ------------------------------------------------------- one-time backfill

def backfill_supplier_excel_links():
    """
    Fills the Supplier Excel field on invoices created before the field
    existed, from the xlsx already attached to each invoice. Run once:

      bench --site <site> execute \
        taal_inventory.api.purchase_excel_upload.backfill_supplier_excel_links

    Safe to re-run. Skips invoices that already have the field filled.
    Works on submitted invoices too (direct db write, no re-validation).
    """
    if not frappe.get_meta("Purchase Invoice").has_field("custom_supplier_excel"):
        print("Add the 'Supplier Excel' field (custom_supplier_excel) to Purchase Invoice first.")
        return

    files = frappe.get_all(
        "File",
        filters={"attached_to_doctype": "Purchase Invoice", "file_url": ["like", "%.xlsx"]},
        fields=["attached_to_name", "file_url"],
        order_by="creation asc",
    )
    linked = 0
    for row in files:
        if not row.attached_to_name:
            continue
        if frappe.db.get_value("Purchase Invoice", row.attached_to_name, "custom_supplier_excel"):
            continue
        frappe.db.set_value("Purchase Invoice", row.attached_to_name,
                            "custom_supplier_excel", row.file_url, update_modified=False)
        linked += 1
    frappe.db.commit()
    print(f"Linked {linked} invoice(s) to their supplier Excel.")


# -------------------------------------------------------------- entry point

@frappe.whitelist()
def preview_supplier_excel(excel_file=None, company=None):
    """
    Called the moment the Excel is attached in the dialog. READ-ONLY: parses
    the file and returns what the dialog should pre-fill. Writes nothing.
    """
    if not frappe.has_permission("Purchase Invoice", "create"):
        frappe.throw(_("You do not have permission to create Purchase Invoices."),
                     frappe.PermissionError)
    if not excel_file:
        frappe.throw(_("No Excel file was attached."))

    file_doc = frappe.get_doc("File", {"file_url": excel_file})
    data = parse_workbook(file_doc.get_content())

    supplier = find_supplier(data)
    template = pick_tax_template(company, data) if company else None
    bill_no = data["ref_no"] or data["invoice_no"]
    bill_date = data["ref_date"] or data["invoice_date"]

    return {
        "supplier": supplier,
        "supplier_name_on_file": data["supplier_name"],
        "supplier_gstin": data["supplier_gstin"],
        "taxes_and_charges": template,
        "bill_no": bill_no,
        "bill_date": str(bill_date) if bill_date else None,
        "posting_date": str(bill_date) if bill_date else None,
        "item_count": len(data["items"]),
        "grand_total": data["grand_total"],
    }


@frappe.whitelist()
def create_purchase_invoice(excel_file=None, company=None, supplier=None,
                            taxes_and_charges=None, set_warehouse=None,
                            posting_date=None, bill_no=None, bill_date=None,
                            create_supplier=0):
    """Called by the Create from Supplier Excel dialog on Purchase Invoice."""
    if not frappe.has_permission("Purchase Invoice", "create"):
        frappe.throw(_("You do not have permission to create Purchase Invoices."),
                     frappe.PermissionError)

    if not excel_file:
        frappe.throw(_("No Excel file was attached."))
    if not company:
        frappe.throw(_("Company is required."))

    args = frappe._dict(
        excel_file=excel_file,
        company=company,
        supplier=supplier or "",
        taxes_and_charges=taxes_and_charges,
        set_warehouse=set_warehouse,
        posting_date=getdate(posting_date) if posting_date else None,
        bill_no=clean(bill_no),
        bill_date=getdate(bill_date) if bill_date else None,
        create_supplier=cint(create_supplier),
    )

    file_doc = frappe.get_doc("File", {"file_url": excel_file})
    content = file_doc.get_content()

    try:
        data = parse_workbook(content)
        errors = validate_parsed(args, data)
        if errors:
            frappe.throw("<br>".join(errors), title=_("Cannot create Purchase Invoice"))

        invoice = build_invoice(args, data)
        summary = verify_invoice(invoice, data)
        _attach_file_to_invoice(file_doc, invoice)

        return {"purchase_invoice": invoice.name, "summary": summary}

    except frappe.ValidationError:
        raise
    except Exception:
        frappe.log_error(frappe.get_traceback(), "Purchase Excel Upload")
        raise